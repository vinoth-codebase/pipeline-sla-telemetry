"""Tests for the CLI, the MS Office reporting engines and the dashboard views."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from click.testing import CliRunner
from openpyxl import load_workbook
from pptx import Presentation

from pipeline_sla.cli import main
from pipeline_sla.dashboard import views
from pipeline_sla.reporting.excel_exporter import ExcelExporter
from pipeline_sla.reporting.pptx_exporter import PptxExporter


@pytest.fixture()
def cli_db(tmp_path, settings) -> Path:
    """A small on-disk warehouse produced by the CLI itself."""
    db = tmp_path / "cli.duckdb"
    result = CliRunner().invoke(main, ["generate", "--days", "40", "--output", str(db)])
    assert result.exit_code == 0, result.output
    return db


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def test_cli_version():
    result = CliRunner().invoke(main, ["--version"])
    assert result.exit_code == 0
    assert "1.0.0" in result.output


def test_cli_help_lists_commands():
    result = CliRunner().invoke(main, ["--help"])
    assert result.exit_code == 0
    for command in ("generate", "analyze", "forecast", "lineage", "export", "dashboard"):
        assert command in result.output


def test_cli_generate(tmp_path):
    db = tmp_path / "gen.duckdb"
    result = CliRunner().invoke(main, ["generate", "--days", "20", "--seed", "3", "--output", str(db)])
    assert result.exit_code == 0
    assert db.exists()
    assert "dag_runs" in result.output


def test_cli_analyze(cli_db):
    result = CliRunner().invoke(main, ["analyze", "--all", "--db", str(cli_db)])
    assert result.exit_code == 0
    assert "Reliability Scorecard" in result.output


def test_cli_analyze_single_dag(cli_db):
    result = CliRunner().invoke(main, ["analyze", "--dag", "crm_customer_360", "--db", str(cli_db)])
    assert result.exit_code == 0
    assert "crm_customer_360" in result.output


def test_cli_analyze_empty_warehouse(tmp_path):
    db = tmp_path / "empty.duckdb"
    from pipeline_sla.storage.warehouse import TelemetryWarehouse

    TelemetryWarehouse(db).close()
    result = CliRunner().invoke(main, ["analyze", "--all", "--db", str(db)])
    assert result.exit_code == 0
    assert "No telemetry" in result.output


def test_cli_forecast(cli_db):
    result = CliRunner().invoke(main, [
        "forecast", "--dag", "gl_financial_reconciliation", "--task", "reconcile_subledgers",
        "--elapsed", "600", "--cutoff", "30", "--db", str(cli_db),
    ])
    assert result.exit_code == 0
    assert "Breach probability" in result.output


def test_cli_lineage_ranking():
    result = CliRunner().invoke(main, ["lineage", "--top", "3"])
    assert result.exit_code == 0
    assert "Blast Radius" in result.output


def test_cli_lineage_single_task():
    result = CliRunner().invoke(main, ["lineage", "--task", "load_sales_fact"])
    assert result.exit_code == 0
    assert "load_sales_fact" in result.output


def test_cli_export(cli_db, tmp_path):
    out = tmp_path / "reports"
    result = CliRunner().invoke(main, ["export", "--excel", "--pptx", "--outdir", str(out), "--db", str(cli_db)])
    assert result.exit_code == 0
    assert (out / "weekly_pipeline_health.xlsx").exists()
    assert (out / "monday_operations_review.pptx").exists()


# --------------------------------------------------------------------------- #
# Excel exporter
# --------------------------------------------------------------------------- #
def test_excel_export_structure(warehouse, tmp_path):
    path = ExcelExporter().export(warehouse, tmp_path / "health.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Executive Overview", "SLA Breach Log", "Task Percentiles"]
    ws = wb["Executive Overview"]
    assert ws["A1"].value.startswith("Weekly Pipeline Health")
    assert ws.conditional_formatting


def test_excel_export_empty_warehouse(empty_warehouse, tmp_path):
    path = ExcelExporter().export(empty_warehouse, tmp_path / "empty.xlsx")
    wb = load_workbook(path)
    assert wb["Executive Overview"]["A5"].value == "No telemetry available."


def test_excel_breach_log_content(warehouse, tmp_path):
    path = ExcelExporter().export(warehouse, tmp_path / "health.xlsx")
    wb = load_workbook(path)
    ws = wb["SLA Breach Log"]
    # Either breaches were logged, or the green all-clear message is present.
    assert ws["A4"].value is not None


def test_excel_percentiles_tab(warehouse, tmp_path):
    path = ExcelExporter().export(warehouse, tmp_path / "health.xlsx")
    wb = load_workbook(path)
    ws = wb["Task Percentiles"]
    headers = [ws.cell(row=4, column=c).value for c in range(1, 9)]
    assert "P95 Seconds" in headers


# --------------------------------------------------------------------------- #
# PowerPoint exporter
# --------------------------------------------------------------------------- #
def test_pptx_export_three_slides(warehouse, tmp_path):
    path = PptxExporter().export(warehouse, tmp_path / "review.pptx")
    prs = Presentation(path)
    assert len(prs.slides._sldIdLst) == 3


def test_pptx_export_empty_warehouse(empty_warehouse, tmp_path):
    path = PptxExporter().export(empty_warehouse, tmp_path / "empty.pptx")
    prs = Presentation(path)
    assert len(prs.slides._sldIdLst) == 3


def test_pptx_custom_title(warehouse, tmp_path):
    path = PptxExporter().export(warehouse, tmp_path / "review.pptx", title="Custom Review")
    prs = Presentation(path)
    texts = [shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert any("Custom Review" in t for t in texts)


def test_pptx_bottlenecks_and_anomalies(warehouse):
    exporter = PptxExporter()
    bottlenecks = exporter._top_bottlenecks(warehouse, top_n=3)
    assert len(bottlenecks) <= 3
    assert "p95_seconds" in bottlenecks.columns
    anomalies = exporter._volume_anomalies(warehouse)
    assert not anomalies.empty


# --------------------------------------------------------------------------- #
# Dashboard views
# --------------------------------------------------------------------------- #
def test_estate_kpis(warehouse, empty_warehouse):
    kpis = views.estate_kpis(warehouse)
    assert kpis["pipelines"] == 6
    assert kpis["total_runs"] > 0
    assert views.estate_kpis(empty_warehouse)["total_runs"] == 0


def test_kpi_overview(warehouse):
    assert len(views.kpi_overview(warehouse)) == 6


def test_parse_cutoff(settings):
    cutoff = views.parse_cutoff("07:00", settings.sla.timezone)
    assert (cutoff.hour, cutoff.minute) == (7, 0)


def test_morning_dispatch(warehouse, settings, empty_warehouse):
    dispatch = views.morning_dispatch(warehouse, settings)
    assert not dispatch.empty
    assert {"dag_id", "status", "colour", "utilisation_pct"} <= set(dispatch.columns)
    assert dispatch.attrs["cutoff"] == "07:00"
    assert views.morning_dispatch(empty_warehouse, settings).empty


def test_weibull_probability_figure(warehouse):
    durations = warehouse.get_task_durations("gl_financial_reconciliation", "reconcile_subledgers")
    fig = views.weibull_probability_figure(durations, elapsed_seconds=600, cutoff_seconds=1800)
    assert len(fig.data) >= 2
    assert "P(breach)" in fig.layout.title.text


def test_weibull_probability_figure_insufficient_data():
    fig = views.weibull_probability_figure([1.0, 2.0], elapsed_seconds=1, cutoff_seconds=10)
    assert "Insufficient" in fig.layout.annotations[0].text


def test_volume_and_cusum_figures(warehouse):
    series = warehouse.get_volume_series("sales_ingestion_etl", "load_sales_fact")
    vol = views.volume_anomaly_figure(series)
    cus = views.cusum_figure(series)
    assert len(vol.data) >= 2
    assert len(cus.data) >= 2


def test_volume_figure_empty():
    fig = views.volume_anomaly_figure(pd.DataFrame())
    assert "No volume data" in fig.layout.annotations[0].text


def test_lineage_figure(lineage_graph):
    fig = views.lineage_figure(lineage_graph, highlight_node="reconcile_subledgers")
    assert len(fig.data) >= 3


def test_lineage_figure_empty_graph():
    from pipeline_sla.lineage.dag_graph import LineageGraph

    fig = views.lineage_figure(LineageGraph())
    assert "Empty lineage graph" in fig.layout.annotations[0].text


def test_blast_radius_table(lineage_graph, settings):
    table = views.blast_radius_table(lineage_graph, settings, top_n=5)
    assert len(table) <= 5
    assert "blast_radius_score" in table.columns


def test_cli_dashboard_importable():
    # `dashboard` shells out to streamlit; only assert the command is wired.
    command = main.commands["dashboard"]
    assert command.params[0].name == "port"
