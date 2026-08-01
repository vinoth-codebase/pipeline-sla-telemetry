"""Excel operations ledger.

Produces a multi-tab, pre-styled workbook (``weekly_pipeline_health.xlsx``) that
an operations lead can open in the **VS Code Excel Viewer** extension and read
at a glance:

* **Executive Overview** -- per-pipeline KPI scorecard with an SLA-attainment
  heatmap (soft green > 99%, amber 95-99%, red < 95%).
* **SLA Breach Log** -- every run that overran its envelope, worst first.
* **Task Percentiles** -- P50/P85/P95 completion boundaries per task.

Everything is written with ``openpyxl`` so conditional-formatting rules,
frozen panes, number formats and column widths survive the round-trip.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pipeline_sla.analytics.metrics import operational_summary
from pipeline_sla.analytics.weibull_forecaster import WeibullForecaster

logger = logging.getLogger(__name__)

# ---- palette ---------------------------------------------------------------
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(color="1F3864", bold=True, size=16)
_SUBTITLE_FONT = Font(color="595959", italic=True, size=10)
_GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
_AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")
_GREEN_FONT = Font(color="006100")
_AMBER_FONT = Font(color="9C6500")
_RED_FONT = Font(color="9C0006")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class ExcelExporter:
    """Builds the weekly pipeline health workbook."""

    def __init__(self, forecaster: WeibullForecaster | None = None) -> None:
        self.forecaster = forecaster or WeibullForecaster()

    # ------------------------------------------------------------------ #
    def export(
        self,
        warehouse: Any,
        output_path: str | Path,
        *,
        start_date: str | None = None,
        end_date: str | None = None,
    ) -> Path:
        """Write the three-tab workbook and return the resolved path."""
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        overview = operational_summary(warehouse, start_date=start_date, end_date=end_date)
        breaches = self._build_breach_log(warehouse, start_date, end_date)
        percentiles = self._build_task_percentiles(warehouse)

        wb = Workbook()
        wb.remove(wb.active)
        self._write_overview(wb.create_sheet("Executive Overview"), overview, start_date, end_date)
        self._write_breaches(wb.create_sheet("SLA Breach Log"), breaches)
        self._write_percentiles(wb.create_sheet("Task Percentiles"), percentiles)
        wb.save(out)
        logger.info("Excel ledger written to %s", out)
        return out

    # ------------------------------------------------------------------ #
    def _build_breach_log(
        self, warehouse: Any, start_date: str | None, end_date: str | None
    ) -> pd.DataFrame:
        runs = warehouse.get_dag_runs()
        if runs.empty:
            return pd.DataFrame()
        runs = runs.copy()
        runs["sla_seconds"] = pd.to_numeric(runs["sla_minutes"], errors="coerce") * 60.0
        runs["overrun_seconds"] = pd.to_numeric(runs["duration_seconds"], errors="coerce") - runs["sla_seconds"]
        breaches = runs[runs["overrun_seconds"] > 0].copy()
        if start_date:
            breaches = breaches[pd.to_datetime(breaches["execution_date"]) >= pd.Timestamp(start_date)]
        if end_date:
            breaches = breaches[pd.to_datetime(breaches["execution_date"]) <= pd.Timestamp(end_date)]
        breaches = breaches.sort_values("overrun_seconds", ascending=False)
        breaches["execution_date"] = pd.to_datetime(breaches["execution_date"]).dt.strftime("%Y-%m-%d")
        cols = ["dag_id", "run_id", "execution_date", "state", "duration_seconds",
                "sla_seconds", "overrun_seconds", "business_tower", "criticality_tier"]
        return breaches[[c for c in cols if c in breaches.columns]].reset_index(drop=True)

    def _build_task_percentiles(self, warehouse: Any) -> pd.DataFrame:
        tasks = warehouse.get_task_instances()
        if tasks.empty:
            return pd.DataFrame()
        tasks = tasks[tasks["state"] == "success"]
        records: list[dict[str, Any]] = []
        for (dag_id, task_id), group in tasks.groupby(["dag_id", "task_id"]):
            durations = pd.to_numeric(group["duration_seconds"], errors="coerce").dropna()
            if durations.empty:
                continue
            records.append(
                {
                    "dag_id": dag_id,
                    "task_id": task_id,
                    "runs": int(len(durations)),
                    "mean_seconds": round(float(durations.mean()), 1),
                    "p50_seconds": round(float(durations.quantile(0.50)), 1),
                    "p85_seconds": round(float(durations.quantile(0.85)), 1),
                    "p95_seconds": round(float(durations.quantile(0.95)), 1),
                    "max_seconds": round(float(durations.max()), 1),
                }
            )
        return pd.DataFrame(records).sort_values(["dag_id", "task_id"]).reset_index(drop=True)

    # ------------------------------------------------------------------ #
    # Sheet writers
    # ------------------------------------------------------------------ #
    def _write_overview(
        self, ws: Worksheet, overview: pd.DataFrame, start_date: str | None, end_date: str | None
    ) -> None:
        ws.sheet_view.showGridLines = False
        ws["A1"] = "Weekly Pipeline Health - Executive Overview"
        ws["A1"].font = _TITLE_FONT
        period = f"Reporting window: {start_date or 'all history'} to {end_date or 'latest'}"
        ws["A2"] = period
        ws["A2"].font = _SUBTITLE_FONT

        headers = [
            "DAG ID", "Total Runs", "Successful", "Failed", "MTTR (s)",
            "MTBF (h)", "First-Run Success %", "SLA Attainment %",
            "Avg Duration (s)", "P95 Duration (s)", "Rows Loaded",
        ]
        start_row = 4
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER

        if overview.empty:
            ws.cell(row=start_row + 1, column=1, value="No telemetry available.")
            return

        for r, (_, row) in enumerate(overview.iterrows(), start=start_row + 1):
            values = [
                row.get("dag_id"), row.get("total_runs"), row.get("successful_runs"),
                row.get("failed_runs"), row.get("mttr_seconds"), row.get("mtbf_hours"),
                row.get("first_run_success_rate"), row.get("sla_attainment"),
                row.get("avg_duration_seconds"), row.get("p95_duration_seconds"),
                row.get("total_rows_loaded"),
            ]
            for c, value in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.border = _BORDER
                if c in {3, 4, 11}:
                    cell.number_format = "#,##0"
                elif c in {5, 6, 9, 10}:
                    cell.number_format = "#,##0.0"
                elif c in {7, 8}:
                    cell.number_format = '0.00"%"'

        last_row = start_row + len(overview)
        # SLA attainment heatmap: green > 99, amber 95-99, red < 95.
        sla_range = f"H{start_row + 1}:H{last_row}"
        ws.conditional_formatting.add(
            sla_range, CellIsRule(operator="greaterThan", formula=["99"], fill=_GREEN_FILL, font=_GREEN_FONT)
        )
        ws.conditional_formatting.add(
            sla_range,
            CellIsRule(operator="between", formula=["95", "99"], fill=_AMBER_FILL, font=_AMBER_FONT),
        )
        ws.conditional_formatting.add(
            sla_range, CellIsRule(operator="lessThan", formula=["95"], fill=_RED_FILL, font=_RED_FONT)
        )
        # MTTR colour scale adds a second, gradient read of recovery speed.
        ws.conditional_formatting.add(
            f"E{start_row + 1}:E{last_row}",
            ColorScaleRule(start_type="min", start_color="C6EFCE", end_type="max", end_color="FFC7CE"),
        )

        ws.freeze_panes = f"A{start_row + 1}"
        self._autosize(ws, headers, last_row)
        ws.column_dimensions["A"].width = 30

    def _write_breaches(self, ws: Worksheet, breaches: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        ws["A1"] = "SLA Breach Log - Runs Overrunning Their Envelope"
        ws["A1"].font = _TITLE_FONT
        ws["A2"] = f"{len(breaches)} breach(es) detected in the reporting window."
        ws["A2"].font = _SUBTITLE_FONT

        if breaches.empty:
            ws["A4"] = "No SLA breaches recorded. Green across the estate."
            ws["A4"].font = Font(color="006100", bold=True)
            return

        headers = [c.replace("_", " ").title() for c in breaches.columns]
        start_row = 4
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER

        for r, (_, row) in enumerate(breaches.iterrows(), start=start_row + 1):
            for c, col_name in enumerate(breaches.columns, start=1):
                value = row[col_name]
                cell = ws.cell(row=r, column=c, value=value)
                cell.border = _BORDER
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "#,##0.0"

        last_row = start_row + len(breaches)
        overrun_col = get_column_letter(list(breaches.columns).index("overrun_seconds") + 1)
        ws.conditional_formatting.add(
            f"{overrun_col}{start_row + 1}:{overrun_col}{last_row}",
            ColorScaleRule(start_type="min", start_color="FFEB9C", end_type="max", end_color="FFC7CE"),
        )
        ws.freeze_panes = f"A{start_row + 1}"
        self._autosize(ws, headers, last_row)

    def _write_percentiles(self, ws: Worksheet, percentiles: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        ws["A1"] = "Task Percentiles - Completion Boundaries"
        ws["A1"].font = _TITLE_FONT
        ws["A2"] = "P50 / P85 / P95 historical task runtimes (seconds) used to seed the Weibull fit."
        ws["A2"].font = _SUBTITLE_FONT

        if percentiles.empty:
            ws["A4"] = "No task history available."
            return

        headers = [c.replace("_", " ").title() for c in percentiles.columns]
        start_row = 4
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _BORDER

        for r, (_, row) in enumerate(percentiles.iterrows(), start=start_row + 1):
            for c, col_name in enumerate(percentiles.columns, start=1):
                value = row[col_name]
                cell = ws.cell(row=r, column=c, value=value)
                cell.border = _BORDER
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "#,##0.0"

        last_row = start_row + len(percentiles)
        p95_col = get_column_letter(list(percentiles.columns).index("p95_seconds") + 1)
        ws.conditional_formatting.add(
            f"{p95_col}{start_row + 1}:{p95_col}{last_row}",
            ColorScaleRule(start_type="min", start_color="C6EFCE", end_type="max", end_color="FFC7CE"),
        )
        ws.freeze_panes = f"A{start_row + 1}"
        self._autosize(ws, headers, last_row)

    # ------------------------------------------------------------------ #
    @staticmethod
    def _autosize(ws: Worksheet, headers: list[str], last_row: int) -> None:
        for col, header in enumerate(headers, start=1):
            letter = get_column_letter(col)
            width = max(len(str(header)) + 4, 12)
            ws.column_dimensions[letter].width = min(width, 34)
        ws.row_dimensions[4].height = 30
